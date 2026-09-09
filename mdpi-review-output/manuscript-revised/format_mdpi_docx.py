"""Apply the official MDPI Agronomy Word styles to the generated manuscript.

The LaTeX sections remain the content source of truth.  ``make_docx.py`` first
converts them with the official Agronomy template as the reference document;
this script then maps the generated paragraphs to the template's named MDPI
styles, supplies the submission-style front matter, and writes a separate clean
DOCX so the previous manuscript is preserved.
"""

from __future__ import annotations

import re
import zipfile
from copy import deepcopy
from pathlib import Path

from lxml import etree

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import authorship


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "manuscript-revised.docx"
TEMPLATE = HERE.parents[1] / "mdpi-template" / "agronomy-template.docx"
# The single file to submit.  Earlier passes left three near-identical
# variants (..._mdpi, ..._mdpi_simulation, ..._mdpi_simulation_short_captions);
# they are gone -- this is the only styled output.
OUTPUT = HERE / "manuscript-revised-mdpi.docx"


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def remove_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._p.getparent()
    parent.remove(paragraph._p)


def paragraph_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    created = Paragraph(new_p, paragraph._parent)
    if style is not None:
        created.style = style
    if text:
        created.add_run(text)
    return created


def prepend_run(paragraph: Paragraph, text: str, *, bold: bool = False) -> None:
    run = OxmlElement("w:r")
    if bold:
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:b"))
        run.append(rpr)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    ppr = paragraph._p.pPr
    insert_at = 1 if ppr is not None else 0
    paragraph._p.insert(insert_at, run)


def remove_text_prefix(paragraph: Paragraph, prefix: str) -> None:
    """Remove a plain-text prefix without discarding later run formatting."""
    remaining = len(prefix)
    for run in paragraph.runs:
        if remaining == 0:
            break
        if len(run.text) <= remaining:
            remaining -= len(run.text)
            run.text = ""
        else:
            run.text = run.text[remaining:]
            remaining = 0
    if remaining:
        raise RuntimeError(f"Could not remove paragraph prefix: {prefix!r}")


def name_runs(data):
    """(text, superscript, bold) runs for the author-names line.

    With no authors.json the placeholders are written exactly as they were, so
    the file still says out loud what it is missing.
    """
    if data is None:
        return [
            ("[[AUTHOR 1 - FULL NAME REQUIRED]]", "1,*", True),
            (" [[ORCID REQUIRED]]; ", "", False),
            ("[[AUTHOR 2 - FULL NAME REQUIRED]]", "2", True),
            (" [[ORCID REQUIRED]]; [[ADD OR DELETE AUTHORS AS REQUIRED]]", "", False),
        ]
    return authorship.word_front_matter(data)["names"]


def affiliation_lines(data):
    """(bold prefix, rest) for each affiliation line and the correspondence line."""
    if data is None:
        return [
            ("1 ", "[[AFFILIATION 1 REQUIRED]]; [[AUTHOR 1 E-MAIL REQUIRED]]"),
            ("2 ", "[[AFFILIATION 2 REQUIRED]]; [[AUTHOR 2 E-MAIL REQUIRED]]"),
            ("* Correspondence: ", "[[CORRESPONDING AUTHOR E-MAIL REQUIRED]]"),
        ]
    front = authorship.word_front_matter(data)
    return list(front["affiliations"]) + [front["correspondence"]]


def add_author_line(paragraph: Paragraph, data) -> None:
    clear_paragraph(paragraph)
    for text, superscript, bold in name_runs(data):
        run = paragraph.add_run(text)
        run.bold = bold
        if superscript:
            sup = paragraph.add_run(superscript)
            sup.font.superscript = True


def set_labeled_placeholder(paragraph: Paragraph, prefix: str, remainder: str) -> None:
    clear_paragraph(paragraph)
    r = paragraph.add_run(prefix)
    r.bold = True
    paragraph.add_run(remainder)


def _raw_style(paragraph) -> str:
    """The pStyle value as written, not as resolved.

    pandoc marks the front matter with its own Title / Author / AbstractTitle
    styles, which the MDPI template does not define, so python-docx resolves all
    three to Normal.  The reference has to be read straight off the paragraph.
    """
    ppr = paragraph._p.pPr
    if ppr is None or ppr.pStyle is None:
        return ""
    return ppr.pStyle.val or ""


def _by_style(paragraphs, style_id):
    return [p for p in paragraphs if _raw_style(p) == style_id]


def set_front_matter(doc: Document) -> None:
    """Article type, title, authors, affiliations, abstract, keywords, rule.

    Found by style rather than by position: the number of paragraphs pandoc
    emits for \author{} depends on how many authors there are, and this pass
    has to survive the day the real ones are filled in.
    """
    data = authorship.load()
    paragraphs = list(doc.paragraphs)

    title = next(iter(_by_style(paragraphs, "Title")), None)
    author_ps = _by_style(paragraphs, "Author")
    abstract_title = next(iter(_by_style(paragraphs, "AbstractTitle")), None)
    keywords = next((p for p in paragraphs if p.text.strip().startswith("Keywords:")), None)
    if title is None or not author_ps or abstract_title is None or keywords is None:
        raise RuntimeError("Unexpected generated front matter")
    abstract = paragraphs[paragraphs.index(abstract_title) + 1]

    article = title.insert_paragraph_before("Article")
    article.style = doc.styles["MDPI_1.1_article_type"]
    title.style = doc.styles["MDPI_1.2_title"]

    names = author_ps[0]
    names.style = doc.styles["MDPI_1.3_authornames"]
    add_author_line(names, data)

    wanted = affiliation_lines(data)
    slots = author_ps[1:]
    affiliation_style = doc.styles["MDPI_1.6_affiliation"]
    while len(slots) < len(wanted):
        slots.append(paragraph_after(slots[-1] if slots else names,
                                     style=affiliation_style))
    for extra in slots[len(wanted):]:
        remove_paragraph(extra)
    for paragraph, (prefix, rest) in zip(slots, wanted):
        paragraph.style = affiliation_style
        set_labeled_placeholder(paragraph, prefix, rest)

    remove_paragraph(abstract_title)
    abstract.style = doc.styles["MDPI_1.7_abstract"]
    prepend_run(abstract, "Abstract: ", bold=True)

    keywords.style = doc.styles["MDPI_1.8_keywords"]
    line = paragraph_after(keywords, style=doc.styles["MDPI_1.9_line"])
    line.alignment = WD_ALIGN_PARAGRAPH.LEFT


def style_body(doc: Document) -> None:
    h1_no = 0
    h2_no = 0
    table_no = 0
    figure_no = 0
    in_references = False

    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        style_id = (
            ppr.pStyle.val
            if ppr is not None and ppr.pStyle is not None
            else ""
        )
        text = paragraph.text.strip()

        if style_id == "Heading1":
            if text == "References":
                paragraph.style = doc.styles["MDPI_2.1_heading1"]
                in_references = True
            else:
                h1_no += 1
                h2_no = 0
                paragraph.style = doc.styles["MDPI_2.1_heading1"]
                prepend_run(paragraph, f"{h1_no}. ")
            continue

        if in_references:
            numbered_prefix = re.match(r"^\d+\.\s+", text)
            if numbered_prefix:
                remove_text_prefix(paragraph, numbered_prefix.group(0))
            paragraph.style = doc.styles["MDPI_8.1_references"]
            continue

        if style_id == "Heading2":
            h2_no += 1
            paragraph.style = doc.styles["MDPI_2.2_heading2"]
            prepend_run(paragraph, f"{h1_no}.{h2_no}. ")
            continue

        if style_id == "Heading3":
            paragraph.style = doc.styles["MDPI_2.3_heading3"]
            continue

        if style_id == "TableCaption":
            table_no += 1
            paragraph.style = doc.styles["MDPI_4.1_table_caption"]
            prepend_run(paragraph, f"Table {table_no}. ", bold=True)
            continue

        if style_id == "ImageCaption":
            paragraph.style = doc.styles["MDPI_5.1_figure_caption"]
            if text.lower().startswith("graphical abstract"):
                prepend_run(paragraph, "Graphical Abstract. ", bold=True)
                # Drop the duplicate words already at the start while retaining
                # the caption's inline mathematics and emphasis.
                for run in paragraph.runs[1:]:
                    if run.text.lower().startswith("graphical abstract. "):
                        run.text = run.text[len("Graphical abstract. ") :]
                        break
            else:
                figure_no += 1
                prepend_run(paragraph, f"Figure {figure_no}. ", bold=True)
            continue

        if style_id == "CaptionedFigure":
            paragraph.style = doc.styles["MDPI_5.2_figure"]
            continue

        if re.fullmatch(r"\(\d+\)", text):
            paragraph.style = doc.styles["MDPI_3.a_equation_number"]
            continue

        has_display_math = bool(paragraph._p.xpath("./m:oMath | ./m:oMathPara"))
        if has_display_math and not text:
            paragraph.style = doc.styles["MDPI_3.9_equation"]
            continue

        if style_id in {"BodyText", "FirstParagraph"}:
            paragraph.style = doc.styles["MDPI_3.1_text"]


def _template_equation_table():
    """The borderless two-cell equation component from the official template.

    Read straight out of the .docx package: the file's content type is
    ``...wordprocessingml.template.main+xml``, which python-docx refuses to open,
    so the XML is taken with zipfile and handed to python-docx's own parser.
    """
    with zipfile.ZipFile(TEMPLATE) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    for tbl in root.iter(qn("w:tbl")):
        styles = [
            node.get(qn("w:val"))
            for node in tbl.iter(qn("w:pStyle"))
        ]
        if "MDPI39equation" in styles and "MDPI3aequationnumber" in styles:
            return etree.tostring(tbl)
    raise RuntimeError("The Agronomy template has no equation component")


def _cell_paragraph(cell):
    p = cell.find(qn("w:p"))
    if p is None:
        raise RuntimeError("Equation component cell has no paragraph")
    for node in list(p):
        if node.tag != qn("w:pPr"):
            p.remove(node)
    return p


def _set_pstyle(p, style_id: str) -> None:
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    style = ppr.find(qn("w:pStyle"))
    if style is None:
        style = OxmlElement("w:pStyle")
        ppr.insert(0, style)
    style.set(qn("w:val"), style_id)


def _strip_baked_number(math) -> None:
    """Defensive: drop a `(n)` (and the spacing before it) left inside the maths."""
    runs = [r for r in math.iter(qn("m:r"))]
    while runs:
        text = "".join(t.text or "" for t in runs[-1].iter(qn("m:t")))
        if re.fullmatch(r"\(\d+\)", text) or (text and not text.strip("\u2001\u2002\u2003 ")):
            runs[-1].getparent().remove(runs[-1])
            runs.pop()
            continue
        break


def number_equations(doc: Document) -> int:
    """Replace each display-equation paragraph with the MDPI equation component."""
    component = _template_equation_table()
    body = doc.element.body
    equations = [p for p in body.findall(qn("w:p")) if p.find(qn("m:oMathPara")) is not None]
    for number, source in enumerate(equations, start=1):
        table = parse_xml(component)
        cells = table.findall(qn("w:tr") + "/" + qn("w:tc"))
        if len(cells) < 2:
            raise RuntimeError("Equation component is not two-celled")
        eq_p = _cell_paragraph(cells[0])
        _set_pstyle(eq_p, "MDPI39equation")
        for node in list(source):
            if node.tag != qn("w:pPr"):
                eq_p.append(deepcopy(node))
        for math in eq_p.iter(qn("m:oMath")):
            _strip_baked_number(math)

        num_p = _cell_paragraph(cells[-1])
        _set_pstyle(num_p, "MDPI3aequationnumber")
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = f"({number})"
        run.append(node)
        num_p.append(run)

        source.addprevious(table)
        body.remove(source)
    return len(equations)


def add_and_style_back_matter(doc: Document) -> None:
    labels = {
        "Supplementary Materials:",
        "Acknowledgments:",
        "Author Contributions:",
        "Funding:",
        "Institutional Review Board Statement:",
        "Informed Consent Statement:",
        "Data Availability Statement:",
        "Conflicts of Interest:",
    }
    # Supplementary Materials and Acknowledgments used to be inserted here.  They
    # are written by the LaTeX assembler now, like the other six, so that one
    # source says what the back matter contains and this pass only styles it.
    for paragraph in doc.paragraphs:
        if any(paragraph.text.startswith(label) for label in labels):
            paragraph.style = doc.styles["MDPI_6.2_back_matter"]


def enable_field_updates(doc: Document) -> None:
    """Ask Word to refresh PAGE / NUMPAGES when the file is opened.

    The template's header prints "n of N" through fields; without this the
    reviewer sees whatever page count was cached when the file was written.
    """
    settings = doc.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def repeat_header_row(table) -> None:
    """Mark row 1 as a heading row so it repeats across page breaks.

    MDPI requires it and four of the sixteen tables reached the styled file
    without it -- pandoc sets it only for some.
    """
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def style_tables(doc: Document) -> None:
    for table in doc.tables:
        table.style = doc.styles["MDPI_4.1_three_line_table"]
        repeat_header_row(table)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip() == "2-3(lr)4-5 Library":
                        clear_paragraph(paragraph)
                        paragraph.add_run("Library")
                    paragraph.style = doc.styles["MDPI_4.2_table_body"]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def validate(doc: Document, n_equations: int = 0) -> None:
    text = "\n".join(p.text for p in doc.paragraphs)
    checks = {
        "in-silico title": "an in silico study" in text,
        "simulation scope in abstract": "comparative simulation evidence" in text,
        "physical-validation limitation": "There is no physical-greenhouse validation" in text,
        "five numbered sections": all(f"{i}. " in text for i in range(1, 6)),
        "sixteen table captions": sum(
            p.style.style_id == "MDPI41tablecaption" for p in doc.paragraphs
        ) == 16,
        "five numbered figures": sum(
            p.text.startswith("Figure ") for p in doc.paragraphs
        ) == 5,
        "graphical abstract unnumbered": any(
            p.text.startswith("Graphical Abstract.") for p in doc.paragraphs
        ),
        "forty-eight references": sum(
            p.style.style_id == "MDPI81references" for p in doc.paragraphs
        ) == 48,
        "three numbered equations": n_equations == 3
        and sum(
            p.style.style_id == "MDPI3aequationnumber" for p in doc.paragraphs
        ) == 0,  # they live in table cells, not in the body flow
        "equation numbers in table cells": [
            cell.text.strip()
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
            if re.fullmatch(r"\(\d\)", cell.text.strip())
        ] == ["(1)", "(2)", "(3)"],
        "no hand-written number left in the maths": not any(
            re.search(r"\)\s*\(\d\)", p.text) for p in doc.paragraphs
        ),
        "every data table repeats its header row": all(
            t.rows[0]._tr.find(qn("w:trPr")) is not None
            and t.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
            for t in doc.tables
            if t.style is not None and t.style.style_id == "MDPI41threelinetable"
        ),
        "page fields refresh on open": (
            doc.settings.element.find(qn("w:updateFields")) is not None
            and doc.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"
        ),
        "back matter": sum(
            p.style.style_id == "MDPI62backmatter" for p in doc.paragraphs
        ) == 8,
    }
    for label, passed in checks.items():
        print(("OK   " if passed else "FAIL ") + label)
    if not all(checks.values()):
        raise RuntimeError("MDPI formatting validation failed")


def main() -> None:
    doc = Document(SOURCE)
    set_front_matter(doc)
    style_body(doc)
    add_and_style_back_matter(doc)
    style_tables(doc)
    n_equations = number_equations(doc)
    enable_field_updates(doc)
    validate(doc, n_equations)
    doc.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
